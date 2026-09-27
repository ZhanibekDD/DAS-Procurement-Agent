"""Real DB/SMTP/HTTP regressions for lot-bound purchasing and append-only prices."""
import hashlib
import io
import json
from concurrent.futures import ThreadPoolExecutor
from email.message import EmailMessage

import pytest
from openpyxl import Workbook
from pydantic import ValidationError

from procurement.catalog import Catalog, ALIASES, workbook_rows
from procurement.document_viewer import office_html
from procurement.models import CampaignCreate, LotCreate, SupplierCreate
from procurement.procurement_flow import ProcurementFlow, lot_snapshot, validate_message
from procurement.service import ConflictError
from procurement.table_ingest import read_table
from procurement.upload_io import upload_request
from test_launch_workflow import workflow, project, lot_payload, FIXTURES
from test_launch_http import http_boundary
from test_launch_mail import smtp_env
from test_sso_adapter import login, headers, ALICE, BOB
from smtp_capture import CaptureSMTP

FBS=[('ФБС 24.4.6','218'),('ФБС 12.4.6','95'),('ФБС 9.4.6','128')]


def fbs(service):
    pr=project(service)
    data=lot_payload(pr['id']);data['title']='Блоки ФБС'
    data['items']=[{'name':name,'quantity':qty,'unit':'шт'} for name,qty in FBS]
    lot=service.create_lot(LotCreate(**data))
    supplier=service.create_supplier(SupplierCreate(name='ТЕСТ ФБС',email='fbs@example.test',region='Воронеж'))
    return lot,supplier


def campaign(service,lot,supplier):
    flow=ProcurementFlow(service)
    data=CampaignCreate(supplier_ids=[supplier['id']],item_ids=[r['id'] for r in lot['items']])
    p=flow.preview(lot['id'],data)
    data=data.model_copy(update={'snapshot_sha256':p['snapshot_sha256'],'preview_sha256':p['preview_sha256']})
    return p,service.create_campaign(lot['id'],data)


def test_fbs_exact_preview_smtp_and_idempotency(workflow,monkeypatch):
    db,s,w=workflow;lot,supplier=fbs(s);p,c=campaign(s,lot,supplier)
    assert not p['approval_required']
    assert [(i['name'],i['quantity']) for i in p['items']]==FBS
    m=c['messages'][0];s.approve_message(m['id'],'staff-a')
    with CaptureSMTP() as smtp:
        smtp_env(monkeypatch,smtp)
        assert w.send(m['id'],True)['status']=='sent'
        assert w.send(m['id'],True)['duplicate']
        assert len(smtp.messages)==1
        text=smtp.messages[0].get_body(preferencelist=('plain',)).get_content()
        for name,qty in FBS:assert f'{name}: {qty} шт' in text
        assert 'ПБ' not in text
    assert s.get_lot(lot['id'])['status']=='rfq_sent'


@pytest.mark.parametrize('tamper',['quantity','name','lot','body','recipient','attachment'])
def test_changed_snapshot_fail_closed_before_smtp(workflow,monkeypatch,tamper):
    db,s,w=workflow;lot,supplier=fbs(s);_,c=campaign(s,lot,supplier);m=c['messages'][0]
    s.approve_message(m['id'],'staff-a')
    with db.connection() as conn:
        if tamper in {'quantity','name'}:
            column=tamper;value='219' if tamper=='quantity' else 'ПБ 63.12'
            conn.execute(f'UPDATE lot_items SET {column}=? WHERE id=?',(value,lot['items'][0]['id']))
        elif tamper=='lot':conn.execute("UPDATE lots SET delivery_address='чужой адрес' WHERE id=?",(lot['id'],))
        elif tamper in {'body','recipient'}:conn.execute(f'UPDATE outbox_messages SET {tamper}=? WHERE id=?',('ПБ' if tamper=='body' else 'other@example.test',m['id']))
        else:
            doc=s.register_source_document(filename='items.xlsx',content=(FIXTURES/'items.xlsx').read_bytes(),document_type='project_section')
            conn.execute('INSERT INTO outbox_attachments VALUES (?,?,?,?,?)',(m['id'],doc['id'],doc['filename'],doc['sha256'],doc['size_bytes']))
    with CaptureSMTP() as smtp:
        smtp_env(monkeypatch,smtp)
        with pytest.raises(ConflictError):w.send(m['id'],True)
        assert not smtp.messages
    assert not db.all('SELECT * FROM mail_deliveries')


def test_foreign_items_wrong_fbs_quantity_and_stale_preview(workflow):
    db,s,w=workflow;lot,supplier=fbs(s)
    other=s.create_lot(LotCreate(**lot_payload(lot['project_id'])))
    with pytest.raises(ConflictError):ProcurementFlow(s).preview(lot['id'],CampaignCreate(supplier_ids=[supplier['id']],item_ids=[other['items'][0]['id']]))
    p=ProcurementFlow(s).preview(lot['id'],CampaignCreate(supplier_ids=[supplier['id']]))
    with db.connection() as conn:conn.execute("UPDATE suppliers SET email='changed@example.test' WHERE id=?",(supplier['id'],))
    with pytest.raises(ConflictError):s.create_campaign(lot['id'],CampaignCreate(supplier_ids=[supplier['id']],preview_sha256=p['preview_sha256']))
    with db.connection() as conn:conn.execute("UPDATE lot_items SET quantity='219' WHERE id=?",(lot['items'][0]['id'],))
    with pytest.raises(ConflictError):ProcurementFlow(s).preview(lot['id'],CampaignCreate(supplier_ids=[supplier['id']]))
    assert not db.all('SELECT * FROM campaigns')


def test_legacy_rows_retained_explicit_review_backfills_only_draft(workflow):
    db,s,w=workflow;lot,supplier=fbs(s);p,c=campaign(s,lot,supplier)
    before=db.all('SELECT * FROM outbox_messages')
    with db.connection() as conn:
        conn.execute('DELETE FROM rfq_message_snapshots');conn.execute('DELETE FROM rfq_snapshots')
    db.initialize();assert db.all('SELECT * FROM outbox_messages')==before
    with pytest.raises(ConflictError):s.create_campaign(lot['id'],CampaignCreate(supplier_ids=[supplier['id']]))
    _,reviewed=campaign(s,lot,supplier);assert reviewed['id']==c['id']
    assert db.all('SELECT * FROM outbox_messages')==before
    with db.connection() as conn:validate_message(conn,s._outbox_context(conn,c['messages'][0]['id']))
    assert db.one('PRAGMA quick_check')['quick_check']=='ok'


@pytest.mark.parametrize('field', ['item_ids','supplier_ids'])
@pytest.mark.parametrize('value',[True,0,-1,'2'])
def test_immutable_ids_are_strict(field,value):
    with pytest.raises(ValidationError):CampaignCreate(**{'supplier_ids':[1],field:[value]})


def price_csv(supplier='Поставщик A',email='a@example.test',price='112',date='2020-01-01',vat='с НДС',until='2099-01-01'):
    keys=list(ALIASES)
    row=dict(item_name='ФБС 24.4.6',specification='бетон B7.5',category='ФБС',unit='шт',unit_price=price,currency='RUB',vat=vat,delivery='доставка включена',region='Воронежская область',minimum_batch='10',price_date=date,valid_until=until,supplier_name=supplier,email=email)
    return (';'.join(keys)+'\n'+';'.join(row.get(k,'') for k in keys)+'\n').encode()


def import_price(s,w,raw,filename='price.csv'):
    doc=s.register_source_document(filename=filename,content=raw,document_type='price_list')
    cat=Catalog(s,w);p=cat.price_preview(doc,read_table(raw,filename));assert not p['errors'],p
    return cat,doc,p,cat.apply_prices(p['preview_id'],True)


def test_prices_exact_dedupe_history_median_not_reliability(workflow):
    db,s,w=workflow
    cat,doc,p,r=import_price(s,w,price_csv());assert r['added']==1
    assert cat.apply_prices(p['preview_id'],True)==r
    rating=s.list_suppliers()[0]['rating']
    import_price(s,w,price_csv('Поставщик B','b@example.test','88'))
    prices=cat.prices('ФБС','B7.5');assert len(prices)==2
    assert {p['price_index_pct'] for p in prices}=={12.0,-12.0}
    assert all(p['market_median']=='100' for p in prices)
    import_price(s,w,price_csv(price='124',date='2020-01-02'))
    assert len(s.list_suppliers())==2 and s.list_suppliers()[0]['rating']==rating
    history=cat.prices();assert len(history)==3 and sum(p['current'] for p in history)==2
    assert db.one('SELECT unit_price FROM supplier_catalog_prices WHERE id=1')['unit_price']=='112'
    assert not cat.prices('ПБ')


def test_unknown_basis_does_not_duplicate_current_or_fake_market(workflow):
    db,s,w=workflow
    cat,_,_,_=import_price(s,w,price_csv(vat='',until=''))
    import_price(s,w,price_csv(price='125',date='2020-01-02',vat='',until=''))
    rows=cat.prices();assert sum(r['current'] for r in rows)==1
    assert all(r['market_median'] is None for r in rows)


def test_price_reimport_changed_payload_corrupt_source_and_cross_owner(workflow):
    from procurement.identity import authenticated_actor
    db,s,w=workflow;cat,doc,p,_=import_price(s,w,price_csv())
    other=cat.price_preview(doc,read_table(price_csv(price='999'),'price.csv'))
    with pytest.raises(ConflictError):cat.apply_prices(other['preview_id'],True)
    assert db.one('SELECT count(*) n FROM supplier_catalog_prices')['n']==1
    token=authenticated_actor.set('other-user')
    try:
        with pytest.raises(Exception):cat.apply_prices(p['preview_id'],True)
    finally:authenticated_actor.reset(token)


def test_catalog_xlsx_exact_conflicting_identifiers_and_parallel_import(workflow):
    from procurement.identity import authenticated_actor
    db,s,w=workflow
    b=Workbook();ws=b.active;ws.title='Прайс'
    for line in price_csv().decode().strip().splitlines():ws.append(line.split(';'))
    stream=io.BytesIO();b.save(stream)
    cat,doc,p,_=import_price(s,w,stream.getvalue(),'price.xlsx')
    assert cat.prices()[0]['source_sheet']=='Прайс'
    raw=price_csv(price='100',date='2020-01-02')
    doc=s.register_source_document(filename='new.csv',content=raw,document_type='price_list')
    preview=cat.price_preview(doc,read_table(raw,'new.csv'))
    def apply(_):
        token=authenticated_actor.set('staff-a')
        try:return cat.apply_prices(preview['preview_id'],True)
        finally:authenticated_actor.reset(token)
    with ThreadPoolExecutor(max_workers=8) as pool:results=list(pool.map(apply,range(8)))
    assert all(r==results[0] for r in results) and results[0]['added']==1
    assert db.one('SELECT count(*) n FROM supplier_catalog_prices')['n']==2
    s.create_supplier(SupplierCreate(name='Другая организация',tax_id='7707083893',region='Воронеж'))
    raw=price_csv().decode().replace(';;a@example.test',';7707083893;a@example.test').encode()
    doc=s.register_source_document(filename='conflict.csv',content=raw,document_type='price_list')
    preview=cat.price_preview(doc,read_table(raw,'conflict.csv'))
    with pytest.raises(ConflictError):cat.apply_prices(preview['preview_id'],True)


def test_pdf_catalog_requires_human_review_reuses_strict_validator(workflow):
    from types import SimpleNamespace
    db,s,w=workflow
    doc=s.register_source_document(filename='scan.pdf',content=(FIXTURES/'russian_scan.pdf').read_bytes(),document_type='price_list')
    extracted=SimpleNamespace(items=[],errors=['Нужна ручная проверка'],supplier_region='',supplier_name='',supplier_tax_id='',supplier_email='',supplier_phone='',document_date='',valid_until='')
    cat=Catalog(s,w);p=cat.extracted_price_preview(doc,extracted)
    assert p['requires_review'] and not db.all('SELECT * FROM supplier_catalog_prices')
    row={k:price_csv().decode().strip().splitlines()[1].split(';')[n] for n,k in enumerate(ALIASES)}
    checked=cat.review_pdf(p['preview_id'],[row]);assert not checked['errors']
    assert cat.apply_prices(checked['preview_id'],True)['added']==1
    assert cat.prices()[0]['source_document_id']==doc['id']


def test_price_dates_not_import_order_and_future_prices_not_current(workflow):
    db,s,w=workflow
    cat,_,_,_=import_price(s,w,price_csv(price='200',date='2020-01-02'))
    import_price(s,w,price_csv(price='100',date='2020-01-01'))
    import_price(s,w,price_csv(price='999',date='2099-01-01'))
    rows=cat.prices()
    assert [r['unit_price'] for r in rows if r['current']]==['200']


def test_material_search_before_history_limit_and_russian_normalization(workflow):
    db,s,w=workflow;cat,_,_,_=import_price(s,w,price_csv())
    fields=[r['name'] for r in db.all('PRAGMA table_info(supplier_catalog_prices)') if r['name']!='id']
    previous=db.one('SELECT * FROM supplier_catalog_prices WHERE id=1')
    values=[tuple((n if k=='source_row' else 'Другой материал' if k=='item_name' else 'Другая категория' if k=='category' else previous[k]) for k in fields) for n in range(3,5004)]
    with db.connection() as conn:conn.executemany('INSERT INTO supplier_catalog_prices('+','.join(fields)+') VALUES ('+','.join('?' for _ in fields)+')',values)
    assert len(cat.prices('  фбс   24.4.6 ','бетон b7.5'))==1


def test_fbs_custom_template_wrong_content_blocked(workflow):
    from procurement.models import TemplateUpsert
    db,s,w=workflow;lot,supplier=fbs(s)
    s.upsert_template('wrong',TemplateUpsert(name='Неверный шаблон',subject='ФБС',body='Плиты ПБ: 30 шт'))
    with pytest.raises(ConflictError):ProcurementFlow(s).preview(lot['id'],CampaignCreate(template_code='wrong',supplier_ids=[supplier['id']]))
    with pytest.raises(ConflictError):s.create_campaign(lot['id'],CampaignCreate(template_code='wrong',supplier_ids=[supplier['id']]))
    assert not db.all('SELECT * FROM campaigns')


def workbook():
    b=Workbook();b.active.title='Итоги';b.active.append(['Материал','Количество','Статус','Источник'])
    b.active.append(['ФБС','441','по ведомости','КР лист 14'])
    for name,status in [('Исходные','по проекту'),('Спецификации','предварительно'),('Источники','не определено')]:
        ws=b.create_sheet(name);ws.append(['Материал','Значение','Статус']);ws.append(['001230040500','0',status])
    b['Спецификации'].append(['ФБС 24.4.6',218,'шт','КР лист 14'])
    raw=io.BytesIO();b.save(raw);return raw.getvalue()


def test_workbook_original_provenance_idempotent_and_safe_office_view(workflow):
    db,s,w=workflow;pr=project(s);raw=workbook();rows=workbook_rows(raw,'calc.xlsx')
    assert {r['status'] for r in rows}>={'по ведомости','по проекту','предварительно','не определено'}
    doc=s.register_source_document(filename='calc.xlsx',content=raw,document_type='project_section',project_id=pr['id'])
    cat=Catalog(s,w);assert not cat.import_workbook(doc,rows,True)['duplicate']
    assert cat.import_workbook(doc,rows,True)['duplicate']
    p=cat.portfolio(pr['id']);assert len(p['workbook_rows'])==len(rows)
    assert all(r['document_id']==doc['id'] and r['sha256']==hashlib.sha256(raw).hexdigest() for r in p['workbook_rows'])
    assert w.document_file(doc).path.read_bytes()==raw
    assert '001230040500' in office_html(raw,'calc.xlsx','Исходные')
    assert 'ФБС' in office_html(raw,'calc.xlsx','Спецификации')


def test_real_http_preview_send_inline_policy_deny_and_wrong_lot(http_boundary,monkeypatch):
    client,a,settings,db=http_boundary;u=login(client,a);h=headers(u)
    import procurement.app as app
    s=app.service;lot,supplier=fbs(s);path=f"/api/procurement/lots/{lot['id']}/preview"
    data={'supplier_ids':[supplier['id']],'item_ids':[r['id'] for r in lot['items']]}
    p=client.post(path,headers=h,json=data).json();assert p['lot_id']==lot['id']
    c=client.post(f"/api/lots/{lot['id']}/campaigns",headers=h,json={**data,'snapshot_sha256':p['snapshot_sha256'],'preview_sha256':p['preview_sha256']}).json()
    mid=c['messages'][0]['id'];url=f'/api/procurement/outbox/{mid}/send'
    assert client.post(url,headers=h,json={'lot_id':999,'snapshot_sha256':p['snapshot_sha256'],'confirmed':True}).status_code==409
    with CaptureSMTP() as smtp:
        smtp_env(monkeypatch,smtp)
        r=client.post(url,headers=h,json={'lot_id':lot['id'],'snapshot_sha256':p['snapshot_sha256'],'confirmed':True})
        assert r.status_code==200,r.text
        assert len(smtp.messages)==1
    assert db.one('SELECT actor FROM mail_deliveries')['actor']==ALICE
    assert client.put('/api/procurement/policy',headers=h,json={'required_roles':['staff']}).status_code==403
    with db.connection() as conn:conn.execute("INSERT INTO procurement_policy VALUES(1,NULL,'RUB','[\"staff\"]','test')")
    lot2=s.create_lot(LotCreate(**lot_payload(lot['project_id'])))
    p2,c2=campaign(s,lot2,supplier);mid2=c2['messages'][0]['id']
    assert p2['approval_required']
    assert client.post(f'/api/procurement/outbox/{mid2}/send',headers=h,json={'lot_id':lot2['id'],'snapshot_sha256':p2['snapshot_sha256'],'confirmed':True}).status_code==409
    assert client.post(f'/api/procurement/outbox/{mid2}/approve',headers=h,json={'confirmed':True}).status_code==403
    assert client.post(f'/api/outbox/{mid2}/approve',headers=h,json={'approved_by':'forged admin'}).status_code==409


def test_http_mail_catalog_workbook_views_acl_range(http_boundary):
    client,a,settings,db=http_boundary;u=login(client,a);h=headers(u)
    mail=EmailMessage();mail['From']='untrusted@example.test';mail['To']='test@example.test';mail.set_content('Прайс')
    mail.add_attachment(price_csv(),maintype='text',subtype='csv',filename='price.csv')
    r=client.post('/api/procurement/catalog/incoming-mail',headers=h,files={'file':('mail.eml',mail.as_bytes())})
    assert r.status_code==200,r.text
    p=r.json()['previews'][0]
    assert client.post(f"/api/procurement/catalog/{p['preview_id']}/apply",headers=h,json={'confirmed':True}).status_code==200
    assert client.get('/api/procurement/catalog?q=ФБС').json()[0]['supplier_name']=='Поставщик A'
    pr=client.post('/api/projects',headers=h,json={'name':'ТЕСТ','region':'Воронеж','delivery_address':'Тест'}).json()
    raw=workbook();url=f"/api/procurement/projects/{pr['id']}/workbook"
    p=client.post(url,headers=h,files={'file':('calc.xlsx',raw)}).json()
    assert client.post(url,headers=h,data={'confirmed':'true','expected_sha256':'0'*64},files={'file':('calc.xlsx',raw)}).status_code==409
    r=client.post(url,headers=h,data={'confirmed':'true','expected_sha256':p['sha256']},files={'file':('calc.xlsx',raw)})
    assert r.status_code==200,r.text
    did=r.json()['document_id'];view=f'/api/procurement/documents/{did}/view?sheet=Исходные'
    assert '001230040500' in client.get(view).text and client.head(view).status_code==200
    dl=f'/api/launch/documents/{did}/download'
    assert client.get(dl,headers={'Range':'bytes=0-15'}).content==raw[:16]
    login(client,a,BOB);a.users[BOB]['modules']=[]
    for method in ['GET','HEAD']:
        assert client.request(method,view,headers={'Range':'bytes=0-15'}).status_code==403
    client.cookies.clear()
    assert client.get(view,headers={'X-OpenWebUI-User-Id':ALICE}).status_code==401


@pytest.mark.parametrize('path',['/api/procurement/catalog/preview','/api/procurement/catalog/incoming-mail','/api/procurement/projects/2/workbook'])
def test_new_uploads_share_streaming_body_limit(path):
    assert upload_request({'type':'http','method':'POST','path':path})


def test_docx_preview_escapes_markup_and_keeps_original():
    import zipfile
    xml=b'<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"><w:body><w:p><w:r><w:t>&lt;script&gt;001230040500&lt;/script&gt;</w:t></w:r></w:p></w:body></w:document>'
    raw=io.BytesIO()
    with zipfile.ZipFile(raw,'w') as z:z.writestr('word/document.xml',xml)
    data=raw.getvalue();digest=hashlib.sha256(data).hexdigest()
    text=office_html(data,'doc.docx');assert '<script>' not in text and '&lt;script&gt;001230040500' in text
    assert hashlib.sha256(data).hexdigest()==digest
