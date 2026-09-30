import io
import json
import subprocess
import sys
import threading
import time
from pathlib import Path
from unittest.mock import patch

import pytest
from openpyxl import Workbook

from procurement.incoming_mail import Inbox, extract_draft, download, poll, poll_bounded, network_deadline
from procurement.catalog import Catalog, ALIASES
from procurement.identity import authenticated_actor
from test_launch_workflow import workflow, FIXTURES
from test_incoming_mail import message, ingest, rows_of, admin, Mailbox, mailbox_env
from test_procurement_redesign import price_csv


def spreadsheet(multi=False, formula=False):
    import csv
    values=list(csv.reader(io.StringIO(price_csv().decode()),delimiter=';'))
    book=Workbook();sheet=book.active
    for row in values:sheet.append(row)
    if formula:sheet.cell(2, list(ALIASES).index('unit_price')+1, '=1+1')
    if multi:book.create_sheet('Другой прайс')
    output=io.BytesIO();book.save(output);return output.getvalue()


def test_xlsx_real_bytes_and_same_company_history(workflow,tmp_path):
    db,s,w=workflow
    inbox,_=ingest(workflow,tmp_path,message(spreadsheet(),'offer.xlsx'))
    aid=db.one('SELECT id FROM inbox_attachments')['id']
    assert inbox.detail(aid)['rows'][0]['item_name']=='ФБС 24.4.6'
    with admin():
        p=inbox.prepare(aid,rows_of(inbox,aid),True);Catalog(s,w).apply_prices(p['preview_id'],True)
    _,second=ingest(workflow,tmp_path,message(price_csv(price='120',date='2026-09-30'),'new.csv'),uid=2)
    aid=db.one('SELECT id FROM inbox_attachments WHERE message_id=?',(second,))['id']
    with admin():
        p=inbox.prepare(aid,rows_of(inbox,aid),True);Catalog(s,w).apply_prices(p['preview_id'],True)
    assert len(db.all('SELECT * FROM suppliers'))==1
    assert [r['unit_price'] for r in db.all('SELECT unit_price FROM supplier_catalog_prices ORDER BY id')]==['112','120']


@pytest.mark.parametrize('multi,formula',[(True,False),(False,True)])
def test_incomplete_excel_not_silently_imported(workflow,tmp_path,multi,formula):
    inbox,_=ingest(workflow,tmp_path,message(spreadsheet(multi,formula),'offer.xlsx'))
    aid=inbox.db.one('SELECT id FROM inbox_attachments')['id'];draft=inbox.detail(aid)
    assert draft.get('errors') or draft['rows'][0]['review_warning']
    assert not inbox.db.all('SELECT * FROM supplier_catalog_prices')


def test_recognition_failure_keeps_original_and_safe_error(workflow,tmp_path,monkeypatch):
    monkeypatch.setattr('procurement.incoming_mail.extract_draft',lambda *a:(_ for _ in ()).throw(RuntimeError('PASSWORD raw-mail-secret')))
    inbox,_=ingest(workflow,tmp_path)
    a=inbox.db.one('SELECT * FROM inbox_attachments');assert inbox.path(a).read_bytes()==price_csv()
    assert 'PASSWORD' not in json.dumps(inbox.detail(a['id'])) and inbox.detail(a['id'])['errors']


def test_blocked_mail_size_advances_to_next_uid(workflow,mailbox_env):
    db,s,w=workflow;inbox=Inbox(s,w);mb=Mailbox({1:message(),2:message(subject='Second')})
    original=mb.uid
    def fetch(command,ids,fields):
        if fields=='(UID RFC822.SIZE)':return 'OK',[b'1 (UID 1 RFC822.SIZE 20971521)',f'2 (UID 2 RFC822.SIZE {len(mb.messages[2])})'.encode()]
        return original(command,ids,fields)
    mb.uid=fetch
    assert poll(inbox,lambda:mb)['status']=='ok'
    assert db.one('SELECT status FROM inbox_messages WHERE uid=1')['status']=='error'
    assert db.one('SELECT status FROM inbox_messages WHERE uid=2')['status']=='review'


@pytest.mark.parametrize('tamper',['size','uid','offset'])
def test_streamed_fetch_rejects_wrong_identity_length(tmp_path,tamper):
    mb=Mailbox({1:b'original'});original=mb.uid
    def fetch(*args):
        typ,data=original(*args);header,body=data[0]
        if tamper=='size':body=body[:-1]
        elif tamper=='uid':header=header.replace(b'UID 1',b'UID 2')
        else:header=header.replace(b'<0>',b'<1>')
        return typ,[(header,body),b')']
    mb.uid=fetch
    with pytest.raises(ValueError):download(mb,1,8,tmp_path/'test.eml')


def test_literal_bound_before_allocation(tmp_path):
    mb=Mailbox({1:b'original'})
    def never_read(n):pytest.fail('Oversize literal allocated')
    mb.read=never_read;mb.uid=lambda *_:mb.read(1000000000)
    with pytest.raises(ValueError):download(mb,1,8,tmp_path/'test.eml')
    assert mb.read is never_read


def test_network_deadline_interrupts_without_logging(monkeypatch):
    stopped=threading.Event();mb=Mailbox({});mb.sock.shutdown=lambda *_:stopped.set()
    monkeypatch.setattr('procurement.incoming_mail.NETWORK_SECONDS',0.02)
    with pytest.raises(TimeoutError):
        with network_deadline(mb):assert stopped.wait(0.5)


@pytest.mark.skipif(sys.platform=='win32',reason='Production worker uses Linux flock')
def test_two_workers_share_kernel_lock(workflow,monkeypatch):
    import fcntl, os
    db,s,w=workflow
    with (Path(db.path).parent/'incoming.lock').open('a') as lock:
        fcntl.flock(lock,fcntl.LOCK_EX|fcntl.LOCK_NB)
        result=subprocess.run([sys.executable,'-m','procurement.incoming_mail','--once'],
            env={**os.environ,'PROCUREMENT_DB_PATH':db.path},capture_output=True,text=True,timeout=10)
        assert result.returncode==0 and json.loads(result.stdout)=={'status':'already_running'}


@pytest.mark.parametrize('error',['wrong-password','connection-timeout'])
def test_auth_outage_one_attempt_then_later_recovery(workflow,mailbox_env,error):
    db,s,w=workflow;inbox=Inbox(s,w);attempts=[]
    def offline():attempts.append(1);raise OSError(error+' SENSITIVE')
    assert poll(inbox,offline)['status']=='error' and len(attempts)==1
    assert 'SENSITIVE' not in json.dumps(inbox.listing())
    assert poll(inbox,lambda:Mailbox({1:message()}))['status']=='ok'


def test_transport_failure_has_one_bounded_retry_and_no_duplicate(workflow,mailbox_env):
    from procurement.sent_mail import PreAuthTransportError
    db,s,w=workflow;inbox=Inbox(s,w);attempts=[]
    def connect():
        attempts.append(1)
        if len(attempts)==1:raise PreAuthTransportError()
        return Mailbox({1:message()})
    result=poll_bounded(inbox,connect,pause=lambda _:None)
    assert len(attempts)==2, result
    assert result['status']=='ok' and result['received']==1
    assert db.one('SELECT error,last_uid FROM inbox_state')['error']==''
    assert len(db.all('SELECT id FROM inbox_messages'))==1
    assert 'PASSWORD' not in json.dumps(result)+json.dumps(inbox.listing())


def test_auth_failure_does_not_retry_and_transport_outage_stays_degraded(workflow,mailbox_env):
    import imaplib
    from procurement.sent_mail import PreAuthTransportError
    db,s,w=workflow;inbox=Inbox(s,w);attempts=[]
    def denied():attempts.append(1);raise imaplib.IMAP4.error('PASSWORD invalid')
    result=poll_bounded(inbox,denied,pause=lambda _:pytest.fail('Auth retry'))
    assert result['status']=='error' and result['retryable'] is False and len(attempts)==1
    attempts.clear()
    def outage():attempts.append(1);raise PreAuthTransportError()
    result=poll_bounded(inbox,outage,pause=lambda _:None)
    assert result['status']=='error' and result['retryable'] is True and len(attempts)==2
    assert db.one('SELECT error FROM inbox_state')['error']
    assert 'PASSWORD' not in json.dumps(result)+json.dumps(inbox.listing())


@pytest.mark.parametrize('failure',[TimeoutError,ConnectionResetError])
def test_indeterminate_login_transport_is_not_retried(workflow,mailbox_env,failure):
    db,s,w=workflow;inbox=Inbox(s,w);attempts=[]
    def indeterminate_login():
        attempts.append(1)
        raise failure('LOGIN may have reached provider; PASSWORD redacted')
    result=poll_bounded(inbox,indeterminate_login,pause=lambda _:pytest.fail('Login retry'))
    assert result['status']=='error' and result['retryable'] is False and len(attempts)==1
    assert 'PASSWORD' not in json.dumps(result)+json.dumps(inbox.listing())


def test_authenticated_transport_failure_retries_once(workflow,mailbox_env):
    db,s,w=workflow;inbox=Inbox(s,w);attempts=[]
    def connect():
        attempts.append(1)
        mailbox=Mailbox({1:message()})
        if len(attempts)==1:
            mailbox.select=lambda *_,**__:(_ for _ in ()).throw(ConnectionResetError('PASSWORD transport'))
        return mailbox
    result=poll_bounded(inbox,connect,pause=lambda _:None)
    assert len(attempts)==2, result
    assert result['status']=='ok' and result['received']==1
    assert len(db.all('SELECT id FROM inbox_messages'))==1
    assert 'PASSWORD' not in json.dumps(result)+json.dumps(inbox.listing())


def test_temporary_dns_is_identified_before_login(monkeypatch):
    import socket
    from procurement.sent_mail import PreAuthTransportError, connect_imap
    monkeypatch.setenv('PROCUREMENT_IMAP_HOST','mail.example.org')
    monkeypatch.setenv('PROCUREMENT_IMAP_TLS','ssl')
    def temporary_dns(*args,**kwargs):raise socket.gaierror(socket.EAI_AGAIN,'PASSWORD private DNS')
    monkeypatch.setattr('procurement.sent_mail.imaplib.IMAP4_SSL',temporary_dns)
    with pytest.raises(PreAuthTransportError) as caught:connect_imap()
    assert 'PASSWORD' not in str(caught.value)


def test_raw_mail_html_never_rendered_and_many_parts_fail_closed(workflow,tmp_path):
    from email.message import EmailMessage
    mail=EmailMessage();mail.set_content('<script>steal()</script>',subtype='html')
    for n in range(21):mail.add_attachment(price_csv(),maintype='text',subtype='csv',filename=f'{n}.csv')
    inbox,_=ingest(workflow,tmp_path,mail.as_bytes())
    assert not inbox.db.all('SELECT * FROM inbox_attachments')
    assert inbox.listing()['messages'][0]['status']=='error'
    assert 'steal' not in json.dumps(inbox.listing())


def test_quota_stops_intake_without_advancing_cursor(workflow,mailbox_env,monkeypatch):
    db,s,w=workflow;inbox=Inbox(s,w)
    monkeypatch.setattr('procurement.incoming_mail.MAX_PRIVATE_BYTES',-1)
    assert poll(inbox,lambda:Mailbox({1:message()}))['status']=='error'
    assert db.one('SELECT last_uid FROM inbox_state')['last_uid']==0
    assert not db.all('SELECT * FROM inbox_attachments')


def test_duplicate_file_from_different_email_not_reimported(workflow,tmp_path):
    db,s,w=workflow;inbox,_=ingest(workflow,tmp_path)
    aid=db.one('SELECT id FROM inbox_attachments')['id']
    with admin():
        p=inbox.prepare(aid,rows_of(inbox,aid),True);Catalog(s,w).apply_prices(p['preview_id'],True)
    _,mid=ingest(workflow,tmp_path,message(subject='forwarded same price'),uid=2)
    aid=db.one('SELECT id FROM inbox_attachments WHERE message_id=?',(mid,))['id']
    from procurement.service import ConflictError
    with pytest.raises(ConflictError):inbox.prepare(aid,rows_of(inbox,aid),True)
    assert len(db.all('SELECT * FROM supplier_catalog_prices'))==1
