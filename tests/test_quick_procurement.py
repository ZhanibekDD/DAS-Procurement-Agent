"""One-screen file intake must persist a safe draft, never mail during upload."""
from pathlib import Path
from io import BytesIO

from openpyxl import Workbook
from procurement.models import LotCreate, SupplierCreate
from test_launch_workflow import workflow, lot_payload, project
from test_launch_mail import smtp_env
from smtp_capture import CaptureSMTP

from test_launch_http import http_boundary
from test_sso_adapter import BOB, login, headers

FIXTURES=Path(__file__).parent/'fixtures'


def test_xlsx_upload_creates_one_draft_with_original_and_no_mail(http_boundary):
    client,authority,settings,db=http_boundary
    user=login(client,authority);h=headers(user)
    project=client.post('/api/projects',headers=h,json={
        'name':'Тестовый объект','region':'Воронежская область','delivery_address':'Воронеж, Тестовая 1'}).json()
    workbook=Workbook();sheet=workbook.active
    sheet.append(['Наименование','Количество','Ед. изм.','Характеристики'])
    sheet.append(['ФБС 24.4.6',218,'шт','бетон B7.5'])
    stream=BytesIO();workbook.save(stream);raw=stream.getvalue()
    response=client.post('/api/procurement/quick-intake',headers=h,data={'project_id':project['id']},
        files={'file':('items.xlsx',raw)})
    assert response.status_code==200,response.text
    result=response.json()
    assert result['status']=='draft'
    lot=result['lot'];doc=result['document']
    assert lot['status']=='draft' and lot['items']
    assert [(a['document_id'],a['sha256']) for a in lot['attachments']]==[(doc['id'],doc['sha256'])]
    assert client.get(f"/api/launch/documents/{doc['id']}/download").content==raw
    assert db.one('SELECT count(*) AS n FROM outbox_messages')['n']==0
    assert db.one('SELECT count(*) AS n FROM mail_deliveries')['n']==0


def test_invalid_sheet_is_persisted_for_correction_without_lot(http_boundary):
    client,authority,settings,db=http_boundary
    user=login(client,authority);h=headers(user)
    project=client.post('/api/projects',headers=h,json={
        'name':'Тестовый объект','region':'Воронежская область','delivery_address':'Воронеж, Тестовая 1'}).json()
    response=client.post('/api/procurement/quick-intake',headers=h,data={'project_id':project['id']},
        files={'file':('items.xlsx',(FIXTURES/'items.xlsx').read_bytes())})
    assert response.status_code==200,response.text
    result=response.json()
    assert result['status']=='needs_review' and result['draft']['errors']
    recovered=client.get('/api/procurement/quick-draft').json()
    assert recovered['draft']['preview_id']==result['draft']['preview_id']
    mapping=result['draft']['mapping']
    remap=client.post(f"/api/procurement/quick-draft/{result['draft']['preview_id']}/remap",
        headers=h,json={'mapping':mapping})
    assert remap.status_code==200,remap.text
    assert remap.json()['draft']['errors']
    other=login(client,authority,BOB)
    assert client.get('/api/procurement/quick-draft').json() is None
    assert client.post(f"/api/procurement/quick-draft/{result['draft']['preview_id']}/remap",
        headers=headers(other),json={'mapping':mapping}).status_code==404
    assert db.one('SELECT count(*) AS n FROM lots')['n']==0


def test_pdf_upload_saves_review_only_and_never_guesses_quantity(http_boundary,monkeypatch):
    from procurement import document_analysis
    client,authority,settings,db=http_boundary
    user=login(client,authority);h=headers(user)
    project=client.post('/api/projects',headers=h,json={
        'name':'Тестовый объект','region':'Воронежская область','delivery_address':'Воронеж, Тестовая 1'}).json()
    monkeypatch.setattr(document_analysis,'extract_pdf_page_review',lambda *_:{
        'mode':'ocr','text':'ФБС 24.4.6 218 шт','lines':[
            {'line':1,'text':'ФБС 24.4.6 218 шт','confidence':0.84},
            {'line':2,'text':'неразборчивая строка','confidence':0.32}]})
    response=client.post('/api/procurement/quick-intake',headers=h,data={'project_id':project['id']},
        files={'file':('scan.pdf',(FIXTURES/'russian_scan.pdf').read_bytes())})
    assert response.status_code==200,response.text
    result=response.json()
    assert result['status']=='needs_review' and result['kind']=='pdf'
    assert result['draft']['rows'][0]['quantity']==''
    assert len(result['draft']['lines'])==2
    assert db.one('SELECT count(*) AS n FROM lots')['n']==0
    assert db.one('SELECT count(*) AS n FROM outbox_messages')['n']==0


def test_multipage_pdf_cannot_silently_create_lot_from_first_page(http_boundary,monkeypatch):
    from pypdf import PdfReader, PdfWriter
    from procurement import document_analysis
    client,authority,settings,db=http_boundary
    user=login(client,authority);h=headers(user)
    project=client.post('/api/projects',headers=h,json={
        'name':'Тестовый объект','region':'Воронежская область','delivery_address':'Воронеж, Тестовая 1'}).json()
    reader=PdfReader(str(FIXTURES/'russian_scan.pdf'))
    writer=PdfWriter();writer.add_page(reader.pages[0]);writer.add_page(reader.pages[0])
    stream=BytesIO();writer.write(stream)
    monkeypatch.setattr(document_analysis,'extract_pdf_page_review',lambda *_:{
        'mode':'ocr','text':'ФБС 24.4.6 218 шт','lines':[
            {'line':1,'text':'ФБС 24.4.6 218 шт','confidence':0.84}]})
    response=client.post('/api/procurement/quick-intake',headers=h,data={'project_id':project['id']},
        files={'file':('two-pages.pdf',stream.getvalue())})
    assert response.status_code==200,response.text
    result=response.json()
    assert result['draft']['page_count']==2
    assert 'только первый лист' in result['reason']
    payload={'confirmed':True,'reviewed_line_ids':[1],
        'lot':{'project_id':project['id'],'title':'ФБС','region':'Воронежская область',
               'delivery_address':'Воронеж, Тестовая 1','response_deadline':'2026-10-10',
               'currency':'RUB','attachment_document_ids':[result['document']['id']],
               'items':[{'name':'ФБС 24.4.6','quantity':'218','unit':'шт'}]}}
    create=client.post(f"/api/launch/pdf-review/{result['draft']['preview_id']}/create",headers=h,json=payload)
    assert create.status_code==422 and 'несколько страниц' in create.text
    assert db.one('SELECT count(*) AS n FROM lots')['n']==0


def test_quick_intake_requires_write_access_and_safe_suffix(http_boundary):
    client,authority,settings,db=http_boundary
    user=login(client,authority);h=headers(user)
    project=client.post('/api/projects',headers=h,json={
        'name':'Тестовый объект','region':'Воронежская область','delivery_address':'Воронеж, Тестовая 1'}).json()
    url='/api/procurement/quick-intake';data={'project_id':project['id']}
    assert client.post(url,data=data,files={'file':('a.xlsx',b'bad')}).status_code==403
    assert client.post(url,headers=h,data=data,files={'file':('a.html',b'<html>')}).status_code==422
    assert db.one('SELECT count(*) AS n FROM source_documents')['n']==0


def test_auto_recipients_require_full_material_coverage_region_and_verified_email(workflow):
    db,service,launch=workflow
    pr=project(service)
    payload=lot_payload(pr['id']);payload['title']='Блоки ФБС'
    payload['items']=[{'name':'ФБС 24.4.6','quantity':'218','unit':'шт'},
                      {'name':'ФБС 12.4.6','quantity':'95','unit':'шт'},
                      {'name':'ФБС 9.4.6','quantity':'128','unit':'шт'}]
    lot=service.create_lot(LotCreate(**payload))
    good=service.create_supplier(SupplierCreate(name='ФБС Тест',region=payload['region'],
        email='fbs@example.test',categories=['ФБС'],verified=True))
    wrong=service.create_supplier(SupplierCreate(name='ПБ Тест',region=payload['region'],
        email='pb@example.test',categories=['ПБ'],verified=True))
    unverified=service.create_supplier(SupplierCreate(name='ФБС Без проверки',region=payload['region'],
        email='other@example.test',categories=['ФБС'],verified=False))
    matched={row['id']:row for row in service.match_suppliers(lot['id'])}
    assert matched[good['id']]['auto_select'] is True
    assert matched[wrong['id']]['auto_select'] is False
    assert matched[unverified['id']]['auto_select'] is False


def test_category_alone_does_not_auto_select_for_unverified_characteristics(workflow):
    db,service,launch=workflow
    pr=project(service)
    payload=lot_payload(pr['id']);payload['title']='Кабель 10 кВ'
    payload['items']=[{'name':'Кабель 10 кВ','quantity':'10','unit':'м',
                       'specification':'марка А, сечение 25 мм²'}]
    lot=service.create_lot(LotCreate(**payload))
    supplier=service.create_supplier(SupplierCreate(name='Категория кабель',region=payload['region'],
        email='cable@example.test',categories=['Кабель'],verified=True))
    result={row['id']:row for row in service.match_suppliers(lot['id'])}
    assert result[supplier['id']]['auto_select'] is False


def test_quick_xlsx_to_real_smtp_uses_server_preview_and_exact_attachment(http_boundary,monkeypatch):
    client,authority,settings,db=http_boundary
    user=login(client,authority);h=headers(user)
    project=client.post('/api/projects',headers=h,json={
        'name':'Объект ФБС','region':'Воронежская область','delivery_address':'Воронеж, Тестовая 1'}).json()
    supplier=client.post('/api/suppliers',headers=h,json={
        'name':'ФБС Проверенный','region':'Воронежская область','email':'test@example.test',
        'categories':['ФБС'],'verified':True}).json()
    book=Workbook();sheet=book.active
    sheet.append(['Наименование','Количество','Ед. изм.'])
    sheet.append(['ФБС 24.4.6',218,'шт'])
    stream=BytesIO();book.save(stream);original=stream.getvalue()
    intake=client.post('/api/procurement/quick-intake',headers=h,data={'project_id':project['id']},
        files={'file':('fbs.xlsx',original)})
    assert intake.status_code==200,intake.text
    assert intake.json()['status']=='draft',intake.json()
    lot=intake.json()['lot'];doc=intake.json()['document']
    matches=client.get(f"/api/lots/{lot['id']}/supplier-matches").json()
    assert [s['id'] for s in matches if s['auto_select']]==[supplier['id']]
    request={'supplier_ids':[supplier['id']],'item_ids':[i['id'] for i in lot['items']],
             'channel':'email','template_code':'rfq-email'}
    preview=client.post(f"/api/procurement/lots/{lot['id']}/preview",headers=h,json=request)
    assert preview.status_code==200,preview.text
    p=preview.json();assert p['messages'][0]['attachments'][0]['sha256']==doc['sha256']
    campaign=client.post(f"/api/lots/{lot['id']}/campaigns",headers=h,json={**request,
        'snapshot_sha256':p['snapshot_sha256'],'preview_sha256':p['preview_sha256']})
    assert campaign.status_code==201,campaign.text
    mid=campaign.json()['messages'][0]['id']
    with CaptureSMTP() as smtp:
        smtp_env(monkeypatch,smtp)
        sent=client.post(f'/api/procurement/outbox/{mid}/send',headers=h,json={
            'lot_id':lot['id'],'snapshot_sha256':p['snapshot_sha256'],'confirmed':True})
        assert sent.status_code==200,sent.text
        assert sent.json()['accepted_by_smtp'] is True
        assert len(smtp.messages)==1
        attachment=list(smtp.messages[0].iter_attachments())[0]
        assert attachment.get_filename()=='fbs.xlsx'
        assert attachment.get_payload(decode=True)==original


def test_named_fbs_upload_with_pb_position_cannot_reach_send(http_boundary):
    client,authority,settings,db=http_boundary
    user=login(client,authority);h=headers(user)
    project=client.post('/api/projects',headers=h,json={
        'name':'Объект ФБС','region':'Воронежская область','delivery_address':'Воронеж, Тестовая 1'}).json()
    supplier=client.post('/api/suppliers',headers=h,json={
        'name':'ФБС Проверенный','region':'Воронежская область','email':'test@example.test',
        'categories':['ФБС'],'verified':True}).json()
    book=Workbook();sheet=book.active
    sheet.append(['Наименование','Количество','Ед. изм.'])
    sheet.append(['ПБ 63.12',4,'шт'])
    stream=BytesIO();book.save(stream)
    intake=client.post('/api/procurement/quick-intake',headers=h,data={'project_id':project['id']},
        files={'file':('Блоки ФБС.xlsx',stream.getvalue())})
    assert intake.status_code==200,intake.text
    lot=intake.json()['lot']
    preview=client.post(f"/api/procurement/lots/{lot['id']}/preview",headers=h,json={
        'supplier_ids':[supplier['id']],'item_ids':[item['id'] for item in lot['items']]})
    assert preview.status_code==409
    assert db.one('SELECT count(*) AS n FROM outbox_messages')['n']==0
