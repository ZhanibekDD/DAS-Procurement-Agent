"""Append-only supplier catalog and traceable project workbook snapshots."""
import json
import re
import statistics
from datetime import date
from decimal import Decimal, InvalidOperation

from .db import utcnow
from .identity import trusted_actor
from .models import SupplierCreate
from .service import ConflictError
from .table_ingest import contacts, valid_inn, read_table, mapped, suggested_mapping, safe_upload, validate_xlsx_expansion
from .upload_io import open_payload

SCHEMA='''
CREATE TABLE IF NOT EXISTS supplier_catalog_prices (
 id INTEGER PRIMARY KEY, supplier_id INTEGER NOT NULL REFERENCES suppliers(id),
 source_document_id INTEGER NOT NULL REFERENCES source_documents(id), source_sheet TEXT NOT NULL, source_row INTEGER NOT NULL,
 item_name TEXT NOT NULL, specification TEXT NOT NULL, category TEXT NOT NULL,
 unit TEXT NOT NULL, unit_price TEXT NOT NULL, currency TEXT NOT NULL, vat TEXT NOT NULL,
 delivery TEXT NOT NULL, region TEXT NOT NULL, minimum_batch TEXT NOT NULL,
 price_date TEXT NOT NULL, valid_until TEXT NOT NULL, created_at TEXT NOT NULL,
 UNIQUE(source_document_id,source_sheet,source_row));
CREATE TABLE IF NOT EXISTS project_workbook_rows (
 document_id INTEGER NOT NULL REFERENCES source_documents(id), project_id INTEGER NOT NULL REFERENCES projects(id),
 sheet TEXT NOT NULL, row_number INTEGER NOT NULL, cells_json TEXT NOT NULL, status TEXT NOT NULL,
 PRIMARY KEY(document_id,sheet,row_number));
CREATE TABLE IF NOT EXISTS project_budgets (
 project_id INTEGER NOT NULL REFERENCES projects(id), currency TEXT NOT NULL, amount TEXT NOT NULL,
 actor TEXT NOT NULL, updated_at TEXT NOT NULL, PRIMARY KEY(project_id,currency));
'''

ALIASES={
 'item_name':{'товар','материал','наименование товара','наименование','item_name'},
 'specification':{'характеристики','спецификация','specification'},'category':{'категория','category'},
 'unit':{'ед','ед.','единица','ед. изм.','unit'},'unit_price':{'цена','цена за единицу','unit_price'},
 'currency':{'валюта','currency'},'vat':{'ндс','vat'},'delivery':{'доставка','delivery'},
 'region':{'регион','region'},'minimum_batch':{'минимальная партия','minimum_batch'},
 'price_date':{'дата','дата прайса','price_date'},'valid_until':{'действует до','valid_until'},
 'supplier_name':{'поставщик','supplier_name'},'tax_id':{'инн','tax_id'},'email':{'email','почта'},'phone':{'телефон','phone'}}
STATUSES={'по ведомости','по проекту','предварительно','не определено'}


def number(value):
    try:
        n=Decimal(str(value).replace(' ','').replace(',','.'))
        if not n.is_finite() or n<0:
            raise ValueError()
        return str(n)
    except (ValueError,InvalidOperation):
        raise ValueError('Цена или количество должны быть неотрицательным числом') from None


def normalize(value):
    return ' '.join(str(value).lower().replace('ё','е').split())


def comparable_basis(row):
    vat=normalize(row['vat'])
    known_vat=bool(re.fullmatch(r'(?:с ндс|без ндс|ндс включен|ндс не включен|(?:ндс\s*)?\d{1,2}(?:[.,]\d+)?\s*%)',vat))
    delivery=normalize(row['delivery'])
    return known_vat and bool(delivery) and delivery not in {'не указано','не определено','неизвестно','по согласованию'} and bool(row['valid_until']) and bool(row['minimum_batch'])


def workbook_rows(content,filename):
    from openpyxl import load_workbook
    safe_upload(content,filename,{'.xlsx'})
    validate_xlsx_expansion(content,filename)
    result=[]
    with open_payload(content) as stream:
        book=load_workbook(stream,read_only=True,data_only=False,keep_links=False)
        try:
            for sheet in book:
                for n,row in enumerate(sheet.iter_rows(),1):
                    if n>10000 or len(row)>100:
                        raise ValueError('Слишком большой лист')
                    cells=[str(c.value) if c.value is not None else '' for c in row]
                    if not any(cells):
                        continue
                    if len(result)>=10000 or any(len(c)>8000 for c in cells):
                        raise ValueError('Слишком большая таблица')
                    status=next((c for c in cells if c in STATUSES),None)
                    if status is None:
                        status='предварительно' if any(c in {'допущение','формула','по схемам','сводно'} or c.startswith('=') for c in cells) else 'по проекту' if 'по чертежу' in cells else 'не определено'
                    result.append({'sheet':sheet.title,'row':n,'cells':cells,'status':status})
        finally:
            book.close()
    return result


class Catalog:
    def __init__(self,service,launch):
        self.service,self.launch,self.db=service,launch,service.db

    def price_preview(self,doc,table,mapping=None):
        chosen=mapping if mapping is not None else suggested_mapping(table['headers'],ALIASES)
        rows=[];errors=[]
        for source in table['rows']:
            try:
                values=mapped(source,chosen,table['headers'],set(ALIASES))
                if not values.get('item_name') or not values.get('unit_price'):
                    raise ValueError('Не найдены наименование / цена; проверьте сопоставление')
                values['unit_price']=number(values['unit_price'])
                for k in ALIASES:
                    values.setdefault(k,'')
                if values['tax_id'] and not valid_inn(values['tax_id']):
                    raise ValueError('Неверный ИНН')
                values.update(contacts(values))
                if not values['supplier_name'] or not (values['tax_id'] or values['email'] or values['phone']):
                    raise ValueError('Нужны поставщик и точный ИНН, email или телефон')
                if not values['unit'] or not re.fullmatch('[A-Z]{3}',values['currency']):
                    raise ValueError('Нужны единица и валюта; значения не угадываются')
                if values['minimum_batch']:
                    values['minimum_batch']=number(values['minimum_batch'])
                for field in ('price_date','valid_until'):
                    if values[field]:
                        date.fromisoformat(values[field])
                if not values['price_date'] or not values['region']:
                    raise ValueError('Нужны дата прайса и регион')
                SupplierCreate(name=values['supplier_name'],tax_id=values['tax_id'],email=values['email'],phone=values['phone'],region=values['region'],categories=[values['category']] if values['category'] else [])
                rows.append({'source_row':source['row'],**values})
            except (ValueError,TypeError) as exc:
                errors.append({'row':source['row'],'error':str(exc)})
        return self.launch.save_preview('price_catalog',{'document_id':doc['id'],'rows':rows,'errors':errors,
             'headers':table['headers'],'mapping':chosen,'sheet':table['sheet']})

    def _supplier(self,conn,values):
        # Conflicting exact identifiers are not reconciled by fuzzy name.
        candidates=[]
        for supplier in conn.execute('SELECT * FROM suppliers'):
            if ((values['tax_id'] and values['tax_id']==supplier['tax_id']) or
                (values['email'] and values['email'].lower()==supplier['email'].lower()) or
                (values['phone'] and re.sub(r'\D','',values['phone'])==re.sub(r'\D','',supplier['phone']))):
                candidates.append(dict(supplier))
        if len(candidates)>1:
            raise ConflictError('Идентификаторы прайса связаны с разными поставщиками; автоматическое объединение запрещено')
        if candidates:
            supplier=candidates[0]
            if not supplier['active'] or (values['tax_id'] and supplier['tax_id'] and values['tax_id']!=supplier['tax_id']):
                raise ConflictError('Удалённый поставщик или конфликт ИНН; восстановите/проверьте вручную')
            updates={k:values[k] for k in ('tax_id','email','phone') if values[k] and not supplier[k]}
            if updates:
                conn.execute('UPDATE suppliers SET '+','.join(k+'=?' for k in updates)+' WHERE id=?',(*updates.values(),supplier['id']))
                self.db.audit('updated_from_price','supplier',supplier['id'],details={'fields':list(updates)},conn=conn)
            return supplier['id']
        from .launch_workflow import supplier_values
        data=SupplierCreate(name=values['supplier_name'],tax_id=values['tax_id'],email=values['email'],phone=values['phone'],region=values['region'],categories=[values['category']] if values['category'] else [])
        fields=supplier_values(data)
        sid=conn.execute('INSERT INTO suppliers('+','.join(fields)+',source,created_at) VALUES ('+','.join('?' for _ in range(len(fields)+2))+')',(*fields.values(),'price_catalog',utcnow())).lastrowid
        self.db.audit('created_from_price','supplier',sid,conn=conn)
        return sid

    def apply_prices(self,pid,confirmed):
        if confirmed is not True:
            raise ValueError('Подтвердите проверку строк прайса')
        with self.db.connection() as conn:
            conn.execute('BEGIN IMMEDIATE')
            preview,data=self.launch.preview(conn,pid,'price_catalog')
            if preview['status']=='applied':
                return json.loads(preview['result_json'])
            if preview['status']!='preview':
                raise ConflictError('Предпросмотр уже обработан')
            if not data['rows']:
                raise ValueError('Нет корректных цен для импорта')
            doc=conn.execute('SELECT * FROM source_documents WHERE id=?',(data['document_id'],)).fetchone()
            if not doc:raise ConflictError('Исходный прайс отсутствует')
            self.launch.document_file(dict(doc))
            report={'added':0,'skipped':0,'errors':data['errors']}
            fields=('item_name','specification','category','unit','unit_price','currency','vat','delivery','region','minimum_batch','price_date','valid_until')
            for values in data['rows']:
                previous=conn.execute('SELECT * FROM supplier_catalog_prices WHERE source_document_id=? AND source_sheet=? AND source_row=?',(data['document_id'],data['sheet'],values['source_row'])).fetchone()
                if previous:
                    if any(previous[k]!=values[k] for k in fields):
                        raise ConflictError('Строка этого исходника уже импортирована с другими данными; загрузите новую версию прайса')
                    report['skipped']+=1;continue
                sid=self._supplier(conn,values)
                conn.execute('INSERT INTO supplier_catalog_prices(supplier_id,source_document_id,source_sheet,source_row,'+','.join(fields)+',created_at) VALUES ('+','.join('?' for _ in range(len(fields)+5))+')',
                    (sid,data['document_id'],data['sheet'],values['source_row'],*(values[k] for k in fields),utcnow()))
                report['added']+=1
            conn.execute("UPDATE launch_previews SET status='applied',result_json=? WHERE id=?",(json.dumps(report),pid))
            self.db.audit('price_catalog_import','source_document',data['document_id'],details=report,conn=conn)
        return report

    def prices(self,query='',specification=''):
        with self.db.connection() as conn:
            conn.create_function('catalog_normalize',1,normalize,deterministic=True)
            rows=[dict(r) for r in conn.execute('''SELECT p.*,s.name AS supplier_name,s.rating AS reliability,s.active
                FROM supplier_catalog_prices p JOIN suppliers s ON s.id=p.supplier_id
                WHERE instr(catalog_normalize(p.item_name||' '||p.category),?)>0
                AND instr(catalog_normalize(p.specification),?)>0
                ORDER BY p.price_date DESC,p.id DESC LIMIT 5000''',(normalize(query),normalize(specification)))]
        # One current offer per supplier; older prices remain visible as history.
        current={};seen=set()
        for r in rows:
            key=(r['supplier_id'],normalize(r['item_name']),normalize(r['specification']),normalize(r['unit']),r['currency'],r['vat'],normalize(r['region']),normalize(r['delivery']),r['minimum_batch'])
            if r['price_date']>date.today().isoformat():
                r['current']=False
                continue
            r['current']=key not in seen and bool(r['active']) and (not r['valid_until'] or r['valid_until']>=date.today().isoformat())
            seen.add(key)
            if r['current'] and comparable_basis(r):current[key]=r
        groups={}
        for key,r in current.items():
            groups.setdefault(key[1:],[]).append(Decimal(r['unit_price']))
        for r in rows:
            key=(normalize(r['item_name']),normalize(r['specification']),normalize(r['unit']),r['currency'],r['vat'],normalize(r['region']),normalize(r['delivery']),r['minimum_batch'])
            values=groups.get(key,[])
            median=statistics.median(values) if len(values)>=2 else None
            r['market_median']=str(median) if median else None
            r['price_index_pct']=round(float((Decimal(r['unit_price'])/median-1)*100),2) if median else None
        return rows[:1000]

    def extracted_price_preview(self,doc,result):
        rows=[]
        for n,item in enumerate(result.items,1):
            rows.append({'source_row':n,'item_name':item.item_name,'specification':item.brand,'category':'',
                'unit':item.unit,'unit_price':item.unit_price,'currency':item.currency or '',
                'vat':'с НДС' if item.vat_included is True else 'без НДС' if item.vat_included is False else '',
                'delivery':'','region':result.supplier_region,'minimum_batch':'','price_date':result.document_date or '',
                'valid_until':result.valid_until or '','supplier_name':result.supplier_name,'tax_id':result.supplier_tax_id,
                'email':result.supplier_email,'phone':result.supplier_phone})
        return self.launch.save_preview('price_catalog_pdf',{'document_id':doc['id'],'rows':rows,
            'errors':result.errors,'requires_review':True,'source_filename':doc['filename']})

    def review_pdf(self,pid,rows):
        with self.db.connection() as conn:
            preview,data=self.launch.preview(conn,pid,'price_catalog_pdf')
            if preview['status']!='preview':raise ConflictError('Предпросмотр уже обработан')
        headers=list(ALIASES)
        # Reuse the same strict validation/dedup/import, not another permissive PDF path.
        table={'headers':headers,'sheet':'PDF','rows':[{'row':n,'cells':[r.get(k,'') for k in headers]} for n,r in enumerate(rows,1)]}
        return self.price_preview({'id':data['document_id']},table,{k:n for n,k in enumerate(headers)})

    def import_workbook(self,doc,rows,confirmed):
        if confirmed is not True:
            raise ValueError('Подтвердите источник расчёта и статусы')
        with self.db.connection() as conn:
            conn.execute('BEGIN IMMEDIATE')
            before=conn.execute('SELECT count(*) FROM project_workbook_rows WHERE document_id=?',(doc['id'],)).fetchone()[0]
            for row in rows:
                conn.execute('INSERT OR IGNORE INTO project_workbook_rows VALUES (?,?,?,?,?,?)',
                    (doc['id'],doc['project_id'],row['sheet'],row['row'],json.dumps(row['cells'],ensure_ascii=False),row['status']))
            if not before:self.db.audit('project_workbook_import','project',doc['project_id'],details={'document_id':doc['id'],'rows':len(rows)},conn=conn)
        return {'document_id':doc['id'],'rows':len(rows),'duplicate':bool(before)}

    def portfolio(self,project_id):
        project=self.service.get_project(project_id)
        rows=self.db.all('''SELECT r.*,d.filename,d.sha256 FROM project_workbook_rows r
            JOIN source_documents d ON d.id=r.document_id WHERE r.project_id=? ORDER BY r.document_id,r.sheet,r.row_number''',(project_id,))
        for row in rows:row['cells']=json.loads(row.pop('cells_json'))
        source_materials=[];headers={}
        for row in rows:
            cells=row['cells'];key=(row['document_id'],row['sheet'])
            if 'Обозначение' in cells and 'Кол-во' in cells and 'Ед.' in cells:
                headers[key]={k:cells.index(k) for k in ('Обозначение','Кол-во','Ед.')}
                continue
            h=headers.get(key)
            if not h or max(h.values())>=len(cells) or any(c.upper().startswith('ИТОГО') for c in cells[:2]):continue
            try:quantity=number(cells[h['Кол-во']])
            except ValueError:continue
            if not cells[h['Обозначение']] or cells[h['Ед.']] not in {'шт','м','м²','м³','кг','т'}:continue
            source_materials.append({'name':cells[h['Обозначение']],'quantity':quantity,'unit':cells[h['Ед.']],
                'source_document_id':row['document_id'],'source_reference':f"{row['filename']} · {row['sheet']} · строка {row['row_number']}",
                'source_sheet':row['sheet'],'source_row':row['row_number'],'source_status':row['status']})
        lots=[l for l in self.service.list_lots() if l['project_id']==project_id]
        return {'project':project,'workbook_rows':rows,'documents':self.db.all('SELECT id,filename,sha256 FROM source_documents WHERE project_id=?',(project_id,)),
            'materials':self.db.all('''SELECT i.*,l.title AS lot_title FROM lot_items i JOIN lots l ON l.id=i.lot_id WHERE l.project_id=?''',(project_id,))+source_materials,
            'procurements':lots,'quotes':[q for l in lots for q in self.service.list_quotes(l['id'])],
            'budget':self.db.all('SELECT currency,amount FROM project_budgets WHERE project_id=?',(project_id,))}
